

import os

import re

import hmac

import sqlite3

from pathlib import Path

from datetime import date

from functools import wraps



from flask import Flask, request, jsonify, session, send_file

from flask_cors import CORS



from io import BytesIO

from openpyxl import Workbook

from openpyxl.styles import Font, PatternFill, Alignment

from werkzeug.security import generate_password_hash, check_password_hash



app = Flask(__name__)



# Set SPENDWISE_SECRET_KEY in your environment for deployment.

app.config.update(

    SECRET_KEY=os.environ.get(

        "SPENDWISE_SECRET_KEY",

        "dev-only-change-this-secret-before-deployment"

    ),

    SESSION_COOKIE_HTTPONLY=True,

    SESSION_COOKIE_SAMESITE="Lax",

    SESSION_COOKIE_SECURE=False,  # Set True when using HTTPS.

    PERMANENT_SESSION_LIFETIME=60 * 60 * 8,

)



CORS(

    app,

    origins=[

        "http://127.0.0.1:5500",

        "http://localhost:5500",

    ],

    supports_credentials=True,

)



DATABASE = Path(__file__).resolve().parent / "spendwise.db"





def get_db():

    connection = sqlite3.connect(DATABASE)

    connection.row_factory = sqlite3.Row

    return connection





def initialize_database():

    with get_db() as connection:

        connection.execute("""

            CREATE TABLE IF NOT EXISTS users (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                name TEXT NOT NULL,

                email TEXT NOT NULL UNIQUE,

                password_hash TEXT NOT NULL,

                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP

            )

        """)



        connection.execute("""

            CREATE TABLE IF NOT EXISTS transactions (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                title TEXT NOT NULL,

                amount REAL NOT NULL CHECK(amount > 0),

                type TEXT NOT NULL CHECK(type IN ('income', 'expense')),

                category TEXT NOT NULL,

                date TEXT NOT NULL

            )

        """)



        columns = {

            row["name"]

            for row in connection.execute(

                "PRAGMA table_info(transactions)"

            ).fetchall()

        }



        if "user_id" not in columns:

            connection.execute("""

                ALTER TABLE transactions

                ADD COLUMN user_id INTEGER REFERENCES users(id)

            """)





def login_required(function):

    @wraps(function)

    def wrapper(*args, **kwargs):

        if not session.get("user_id"):

            return jsonify({"error": "Please log in first"}), 401

        return function(*args, **kwargs)

    return wrapper





def valid_transaction(data):

    if not isinstance(data, dict):

        return None, "Valid JSON data is required"



    title = data.get("title")

    amount = data.get("amount")

    transaction_type = data.get("type")

    category = data.get("category")

    transaction_date = data.get("date")



    if not isinstance(title, str) or not title.strip() or len(title) > 80:

        return None, "Enter a name of up to 80 characters"



    if (

        isinstance(amount, bool)

        or not isinstance(amount, (int, float))

        or amount <= 0

        or amount > 1_000_000_000

    ):

        return None, "Enter a valid positive amount"



    if transaction_type not in ("income", "expense"):

        return None, "Type must be income or expense"



    if (

        not isinstance(category, str)

        or not category.strip()

        or len(category) > 50

    ):

        return None, "Enter a valid category"



    if not isinstance(transaction_date, str):

        return None, "Enter a date in YYYY-MM-DD format"



    try:

        if date.fromisoformat(transaction_date).isoformat() != transaction_date:

            raise ValueError

    except ValueError:

        return None, "Enter a valid date in YYYY-MM-DD format"



    return {

        "title": title.strip(),

        "amount": amount,

        "type": transaction_type,

        "category": category.strip(),

        "date": transaction_date,

    }, None





@app.route("/")

def home():

    return jsonify({

        "message": "Welcome to the SpendWise API!",

        "status": "running"

    })





# ---------- AUTHENTICATION ----------



@app.route("/api/auth/register", methods=["POST"])

def register():

    data = request.get_json(silent=True)



    if not isinstance(data, dict):

        return jsonify({"error": "Valid JSON data is required"}), 400



    name = data.get("name")

    email = data.get("email")

    password = data.get("password")



    if not isinstance(name, str) or not name.strip() or len(name.strip()) > 80:

        return jsonify({"error": "Enter your name (up to 80 characters)"}), 400



    if not isinstance(email, str) or len(email) > 254:

        return jsonify({"error": "Enter a valid email address"}), 400



    email = email.strip().lower()



    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):

        return jsonify({"error": "Enter a valid email address"}), 400



    if not isinstance(password, str) or len(password) < 8 or len(password) > 128:

        return jsonify({"error": "Password must be 8–128 characters"}), 400



    password_hash = generate_password_hash(password)



    try:

        with get_db() as connection:

            cursor = connection.execute("""

                INSERT INTO users (name, email, password_hash)

                VALUES (?, ?, ?)

            """, (name.strip(), email, password_hash))



            user_id = cursor.lastrowid



            # Keep the existing transactions by assigning legacy records

            # to the first account created after this upgrade.

            legacy_user_count = connection.execute(

                "SELECT COUNT(*) FROM users"

            ).fetchone()[0]



            if legacy_user_count == 1:

                connection.execute("""

                    UPDATE transactions

                    SET user_id = ?

                    WHERE user_id IS NULL

                """, (user_id,))



    except sqlite3.IntegrityError:

        return jsonify({"error": "An account with this email already exists"}), 409



    session.clear()

    session.permanent = True

    session["user_id"] = user_id



    return jsonify({

        "message": "Account created successfully",

        "user": {"id": user_id, "name": name.strip(), "email": email}

    }), 201





@app.route("/api/auth/login", methods=["POST"])

def login():

    data = request.get_json(silent=True)



    if not isinstance(data, dict):

        return jsonify({"error": "Valid JSON data is required"}), 400



    email = data.get("email")

    password = data.get("password")



    if not isinstance(email, str) or not isinstance(password, str):

        return jsonify({"error": "Email and password are required"}), 400



    with get_db() as connection:

        user = connection.execute("""

            SELECT id, name, email, password_hash

            FROM users

            WHERE email = ?

        """, (email.strip().lower(),)).fetchone()



    if not user or not check_password_hash(user["password_hash"], password):

        return jsonify({"error": "Incorrect email or password"}), 401



    session.clear()

    session.permanent = True

    session["user_id"] = user["id"]



    return jsonify({

        "message": "Login successful",

        "user": {

            "id": user["id"],

            "name": user["name"],

            "email": user["email"]

        }

    })





@app.route("/api/auth/me", methods=["GET"])

@login_required

def current_user():

    with get_db() as connection:

        user = connection.execute("""

            SELECT id, name, email

            FROM users

            WHERE id = ?

        """, (session["user_id"],)).fetchone()



    if user is None:

        session.clear()

        return jsonify({"error": "User account not found"}), 401



    return jsonify({"user": dict(user)})





@app.route("/api/auth/logout", methods=["POST"])

def logout():

    session.clear()

    return jsonify({"message": "Logged out successfully"})





# ---------- TRANSACTIONS ----------



@app.route("/api/transactions", methods=["GET"])

@login_required

def get_transactions():

    with get_db() as connection:

        rows = connection.execute("""

            SELECT id, title, amount, type, category, date

            FROM transactions

            WHERE user_id = ?

            ORDER BY date DESC, id DESC

        """, (session["user_id"],)).fetchall()



    return jsonify([dict(row) for row in rows])





@app.route("/api/transactions", methods=["POST"])

@login_required

def add_transaction():

    data, error = valid_transaction(request.get_json(silent=True))



    if error:

        return jsonify({"error": error}), 400



    with get_db() as connection:

        cursor = connection.execute("""

            INSERT INTO transactions

                (title, amount, type, category, date, user_id)

            VALUES (?, ?, ?, ?, ?, ?)

        """, (

            data["title"], data["amount"], data["type"],

            data["category"], data["date"], session["user_id"]

        ))



        row = connection.execute("""

            SELECT id, title, amount, type, category, date

            FROM transactions

            WHERE id = ? AND user_id = ?

        """, (cursor.lastrowid, session["user_id"])).fetchone()



    return jsonify(dict(row)), 201





@app.route("/api/transactions/<int:transaction_id>", methods=["PUT"])

@login_required

def update_transaction(transaction_id):

    data, error = valid_transaction(request.get_json(silent=True))



    if error:

        return jsonify({"error": error}), 400



    with get_db() as connection:

        cursor = connection.execute("""

            UPDATE transactions

            SET title = ?, amount = ?, type = ?, category = ?, date = ?

            WHERE id = ? AND user_id = ?

        """, (

            data["title"], data["amount"], data["type"],

            data["category"], data["date"],

            transaction_id, session["user_id"]

        ))



        if cursor.rowcount == 0:

            return jsonify({"error": "Transaction not found"}), 404



        row = connection.execute("""

            SELECT id, title, amount, type, category, date

            FROM transactions

            WHERE id = ? AND user_id = ?

        """, (transaction_id, session["user_id"])).fetchone()



    return jsonify(dict(row))





@app.route("/api/transactions/<int:transaction_id>", methods=["DELETE"])

@login_required

def delete_transaction(transaction_id):

    with get_db() as connection:

        cursor = connection.execute("""

            DELETE FROM transactions

            WHERE id = ? AND user_id = ?

        """, (transaction_id, session["user_id"]))



        if cursor.rowcount == 0:

            return jsonify({"error": "Transaction not found"}), 404



    return jsonify({"message": "Transaction deleted successfully"})





# ---------- EXCEL EXPORT ----------

@app.route("/api/transactions/export", methods=["GET"])
@login_required
def export_transactions():
    """Download the logged-in user's transactions as an Excel workbook."""
    with get_db() as connection:
        rows = connection.execute("""
            SELECT title, amount, type, category, date
            FROM transactions
            WHERE user_id = ?
            ORDER BY date DESC, id DESC
        """, (session["user_id"],)).fetchall()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "My Transactions"

    headers = ["Transaction", "Amount (INR)", "Type", "Category", "Date"]
    sheet.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="7057F5")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        sheet.append([
            row["title"],
            row["amount"],
            row["type"].title(),
            row["category"],
            row["date"],
        ])

    for cell in sheet["B"][1:]:
        cell.number_format = '"₹"#,##0.00'

    for column, width in {"A": 28, "B": 18, "C": 15, "D": 20, "E": 16}.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    excel_file = BytesIO()
    workbook.save(excel_file)
    excel_file.seek(0)

    return send_file(
        excel_file,
        as_attachment=True,
        download_name="SpendWise_Transactions.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


initialize_database()



if __name__ == "__main__":

    app.run(debug=True, port=5000)

